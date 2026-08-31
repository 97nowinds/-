from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BASE_DIR = Path(__file__).resolve().parent
OUTPUT_PATH = BASE_DIR / "deliverables" / "摄像头算法与RTSP接入路线图.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "0B2545"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
CALLOUT = "F4F6F9"
MUTED = "5B6573"


def set_run_font(run, size=11, bold=False, color="000000"):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.autofit = False
    table_pr = table._tbl.tblPr
    table_width = sum(widths)
    tbl_w = table_pr.first_child_found_in("w:tblW")
    tbl_w.set(qn("w:w"), str(table_width))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = table_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        table_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    layout = table_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        table_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for index, width in enumerate(widths):
        grid.gridCol_lst[index].set(qn("w:w"), str(width))
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths[index] / 1440)
            tc_w = cell._tc.tcPr.tcW
            tc_w.set(qn("w:w"), str(widths[index]))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)


def add_paragraph(doc, text="", size=11, bold=False, color="000000", after=6, before=0, align=None):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.10
    if align is not None:
        paragraph.alignment = align
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    return paragraph


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.line_spacing = 1.167
    set_run_font(paragraph.add_run(text), size=10.5)


def add_heading(doc, text, level):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.keep_with_next = True
    if level == 1:
        size, color, before, after = 16, BLUE, 16, 8
    else:
        size, color, before, after = 13, BLUE, 12, 6
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.10
    set_run_font(paragraph.add_run(text), size=size, bold=True, color=color)
    return paragraph


def write_cell(cell, text, bold=False, color="000000", size=10):
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.10
    set_run_font(paragraph.add_run(text), size=size, bold=bold, color=color)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    for index, header in enumerate(headers):
        set_cell_shading(table.rows[0].cells[index], LIGHT_BLUE)
        write_cell(table.rows[0].cells[index], header, bold=True, color=INK)
    for row_values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row_values):
            write_cell(cells[index], value)
    return table


def add_callout(doc, text):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, CALLOUT)
    write_cell(cell, text, bold=True, color=INK, size=10.5)
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(2)


def configure_document(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.paragraph_format.space_after = Pt(0)
    set_run_font(header.add_run("智能实验系统 | 摄像头与算法路线图"), size=8.5, color=MUTED)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.space_before = Pt(0)
    set_run_font(footer.add_run("内部技术路线图 | 2026-08-09"), size=8.5, color=MUTED)


def build_document():
    doc = Document()
    configure_document(doc)

    add_paragraph(doc, "技术决策简报", size=10, bold=True, color=MUTED, after=4)
    title = add_paragraph(doc, "实验室人员跟踪算法与 RTSP 接入路线图", size=23, bold=True, color=INK, after=4)
    title.paragraph_format.keep_with_next = True
    add_paragraph(doc, "本周算法 Demo 与 2026-08-31 稳定版交付", size=13, color=MUTED, after=14)

    metadata = add_table(
        doc,
        ["项目", "当前决策", "关键日期"],
        [["综合实验室安全预警系统", "网络摄像机直连电脑服务器，不依赖 NVR", "Demo：本周；稳定版：8 月 31 日"]],
        [2600, 4060, 2700],
    )
    for cell in metadata.rows[1].cells:
        set_cell_shading(cell, LIGHT_GRAY)

    add_heading(doc, "一、结论与边界", 1)
    add_callout(doc, "结论：电脑可以直接拉取网络摄像机 RTSP 码流并运行算法，不需要硬盘录像机。NVR 只在录像存储、集中管理或设备本身没有网络编码能力时才需要。")
    for text in [
        "不以海康私有 SDK 作为算法前置条件：RTSP 取视频，ONVIF/ISAPI 用于发现和配置，SDK 仅用于后续设备控制或事件接入。",
        "现有海康设备没有具体型号，不能仅凭销售表述判定。先读取设备标签或 SADP 扫描结果，再按 RTSP/ONVIF 验收清单决定保留或退换。",
        "若只有 BNC/同轴输出，或只有厂商云平台而没有局域网 RTSP/ONVIF，则不能直接作为本项目算法输入。",
    ]:
        add_bullet(doc, text)

    add_heading(doc, "二、推荐摄像头与网络结构", 1)
    add_paragraph(doc, "推荐型号：海康 DS-2CD1347G2-L，4 mm 固定镜头版本，建议 2 台。", size=11, bold=True, color=INK, after=5)
    add_table(
        doc,
        ["选择项", "理由", "部署要求"],
        [
            ["4 MP + 4 mm", "相对 2 MP 提供更高的人脸与上半身像素余量；约 75° 水平视场更适合实验台定向机位。", "机位与人脸高度平齐；人与镜头距离优先控制在 2.7-6 m。"],
            ["RTSP / ONVIF / ISAPI / SDK", "可由 OpenCV 直接拉流，不被单一厂商应用锁定；后续仍可使用海康 SDK。", "创建独立算法账户；开启 RTSP 与 ONVIF。"],
            ["PoE", "一根网线同时提供数据与供电，减少实验室布线和掉电点。", "802.3af PoE 交换机或注入器，Cat5e 网线。"],
        ],
        [2100, 4400, 2860],
    )
    add_paragraph(doc, "部署链路：网络摄像机 -> PoE 交换机 -> 实验室局域网 -> 电脑服务器 -> Flask 页面 / 算法 / 飞书。", size=10.5, bold=True, color=DARK_BLUE, after=4)
    add_paragraph(doc, "RTSP 地址示例：rtsp://用户名:密码@摄像头IP:554/Streaming/Channels/102。使用 102 子码流做低延迟算法，主码流用于取证或回放。", size=10.5, color="000000", after=7)

    add_heading(doc, "三、现有海康设备验收步骤", 1)
    add_table(
        doc,
        ["步骤", "动作", "通过标准"],
        [
            ["1", "查看机身标签，记录完整型号；确认有 RJ45 网口。", "不是仅 BNC/同轴模拟设备。"],
            ["2", "用 SADP 或路由器 DHCP 列表发现设备 IP，电脑与摄像机置于同一局域网。", "浏览器可访问设备配置页。"],
            ["3", "在 Network / Advanced / Integration Protocol 启用 RTSP 与 ONVIF，创建算法账户。", "保存后能获得 RTSP URL。"],
            ["4", "用 VLC 或项目 probe_rtsp.py 读取子码流。", "连续 30 帧，无认证失败或花屏。"],
            ["5", "把 URL 写入 Windows 环境变量，启动本项目。", "页面状态 running，FPS > 10；断网后自动重连。"],
        ],
        [700, 4600, 4060],
    )

    add_heading(doc, "四、算法路线", 1)
    add_paragraph(doc, "本周 Demo 使用已能运行的算法链路：", size=11, bold=True, color=INK, after=4)
    for text in [
        "人脸检测：OpenCV Haar Cascade 提供正脸候选框。",
        "身份确认：LBPH 本地识别；每位授权人员重新采集 50-80 张双机位样本。",
        "单镜头无正脸跟踪：KCF 上半身相关跟踪。正脸确认后，转头、低头、短时遮挡持续维持 track_id 与身份。",
        "跨镜头换手：利用离开时间、另一镜头的唯一候选和冲突保护继承身份；多人候选时拒绝串号。",
        "器材关联：固定摄像头使用多边形区域标定，以人员框底部锚点确定实验台位置。",
    ]:
        add_bullet(doc, text)
    add_paragraph(doc, "8 月 31 日稳定版升级：YOLO 人体检测 + ByteTrack/BoT-SORT 多目标跟踪 + OSNet/人体 ReID 跨镜头特征排序。时间规则保留为约束，避免仅凭时间造成身份串号。", size=10.5, color=DARK_BLUE, after=7)

    add_heading(doc, "五、排期与验收", 1)
    add_table(
        doc,
        ["时间", "交付", "现场验收"],
        [
            ["8/9-8/10", "确认现有设备型号、IP、RTSP/ONVIF 能力；完成选型。", "型号照片、官方规格、IP 发现记录。"],
            ["8/11-8/12", "电脑直拉 RTSP，完成断线重连和状态诊断。", "连续 30 帧；无需 NVR。"],
            ["8/13-8/16", "双路算法 Demo：识别、无正脸跟踪、跨镜头换手、器材区域联动。", "现场走动演示；保存结果日志。"],
            ["8/17-8/23", "接入人体检测/ReID，完成 5 人以内多人测试。", "侧脸、背对、遮挡、交叉场景。"],
            ["8/24-8/30", "稳定性和性能调优，重采集双机位授权样本。", "8 小时运行、FPS/重连/ID 切换记录。"],
            ["8/31", "最终演示、配置、代码和测试报告归档。", "按验收清单逐项通过。"],
        ],
        [1500, 4700, 3160],
    )

    add_heading(doc, "六、本周可展示与风险", 1)
    add_callout(doc, "本周可展示：电脑作为服务器直接拉两路视频；授权人员注册、姓名识别、无正脸持续跟踪、跨镜头换手、器材区域定位和操作事件时间戳联动。")
    for text in [
        "当前人员库仍有旧的 25 张样本，不能作为最终精度结论；必须重新采集 50-80 张双机位样本。",
        "完全背对、长时间遮挡和多人近距离交叉的稳定身份识别，属于 ReID 阶段验收范围，不应在本周 Demo 中承诺百分之百准确。",
        "RTSP 接入的最终结论依赖现有设备的完整型号、局域网 IP 和管理员授权；这些信息尚未提供。",
    ]:
        add_bullet(doc, text)

    add_heading(doc, "七、官方资料", 1)
    sources = [
        "海康 DS-2CD1347G2-L 官方数据表（2026-07-03）：https://assets.hikvision.com/prd/normal/all/doc/sm000041033/DS-2CD1347G2-L_Datasheet_20260703.pdf",
        "海康 RTSP 地址格式说明（2025-10-30）：https://supportusa.hikvision.com/support/solutions/articles/17000129064-how-do-i-get-my-rtsp-stream-",
        "海康 Device Network SDK 下载页：https://display.hikvision.com/en/support/tools/hitools/clc33f9cd4117c753a/",
    ]
    for source in sources:
        add_paragraph(doc, source, size=9.5, color=MUTED, after=4)

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT_PATH)
    print(OUTPUT_PATH)


if __name__ == "__main__":
    build_document()
