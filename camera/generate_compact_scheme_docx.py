from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BASE_DIR = Path(__file__).resolve().parent
OUT_PATH = BASE_DIR / "deliverables" / "综合实验室安全预警系统_简化方案.docx"

BLUE = "1F4E78"
LIGHT_BLUE = "DCEAF4"
LIGHT_GRAY = "F2F4F7"
PALE_RED = "FDECEC"
TEXT = "202124"
TABLE_WIDTH = 9360


def set_font(run, size=11, bold=None, color=TEXT):
    run.font.name = "Calibri"
    props = run._element.get_or_add_rPr()
    props.rFonts.set(qn("w:ascii"), "Calibri")
    props.rFonts.set(qn("w:hAnsi"), "Calibri")
    props.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    node = OxmlElement("w:shd")
    node.set(qn("w:fill"), fill)
    tc_pr.append(node)


def set_cell_margins(cell, top=80, bottom=80, start=120, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side, value in (("top", top), ("bottom", bottom), ("start", start), ("end", end)):
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        margins.append(node)


def set_table_width(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            tc_pr = cell._tc.get_or_add_tcPr()
            cell_width = OxmlElement("w:tcW")
            cell_width.set(qn("w:w"), str(width))
            cell_width.set(qn("w:type"), "dxa")
            tc_pr.append(cell_width)


def add_table(doc, headers, rows, widths, font_size=9.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for col, header in enumerate(headers):
        cell = table.rows[0].cells[col]
        shade(cell, LIGHT_BLUE)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        para = cell.paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.space_after = Pt(0)
        set_font(para.add_run(header), font_size, True, BLUE)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for col, value in enumerate(values):
            cell = cells[col]
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if row_index % 2:
                shade(cell, "F8FAFC")
            para = cell.paragraphs[0]
            para.paragraph_format.space_after = Pt(0)
            para.paragraph_format.line_spacing = 1.1
            if col == 0:
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_font(para.add_run(value), font_size)
    set_table_width(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_callout(doc, title, text, warning=False):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    shade(cell, PALE_RED if warning else LIGHT_GRAY)
    set_cell_margins(cell, 120, 120, 160, 160)
    para = cell.paragraphs[0]
    para.paragraph_format.space_after = Pt(0)
    set_font(para.add_run(f"{title}  "), 10.5, True, "9B1C1C" if warning else BLUE)
    set_font(para.add_run(text), 10.5)
    set_table_width(table, [TABLE_WIDTH])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_bullets(doc, items):
    for item in items:
        para = doc.add_paragraph(style="List Bullet")
        set_font(para.add_run(item))


def add_numbered(doc, items):
    for item in items:
        para = doc.add_paragraph(style="List Number")
        set_font(para.add_run(item))


def configure(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.18
    for name, size in (("Heading 1", 15), ("Heading 2", 12.5)):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLUE)
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(header.add_run("综合实验室安全预警系统 | 简化方案"), 8.5, color="667085")


def add_flow(doc):
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    values = [
        "传感器\n温湿度 / 烟雾 / 火焰",
        "STM32F103C8T6\n采集、滤波、基础判断",
        "USB 串口\nCP2102 或 CH340\n供电 + 数据",
        "电脑预警平台\n现场画面、事件、飞书",
    ]
    colors = ["E8F4F8", "DCEAF4", "FFF4E5", "E8F4F8"]
    for cell, value, color in zip(table.rows[0].cells, values, colors):
        shade(cell, color)
        set_cell_margins(cell, 180, 180, 100, 100)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        para = cell.paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.space_after = Pt(0)
        set_font(para.add_run(value), 9.5, True if "STM32" in value else None, BLUE if "STM32" in value else TEXT)
    set_table_width(table, [2340, 2340, 2340, 2340])
    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_after = Pt(5)
    set_font(caption.add_run("现场设备通过 USB 线直接接入运行预警平台的电脑；不使用网络模块。"), 9, color="667085")


def build_document():
    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure(doc)
    doc.core_properties.title = "综合实验室安全预警系统简化方案"
    doc.core_properties.subject = "STM32F103C8T6 USB 串口环境监测方案"
    doc.core_properties.author = "项目组"

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(12)
    title.paragraph_format.space_after = Pt(6)
    set_font(title.add_run("综合实验室安全预警系统"), 24, True, BLUE)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(14)
    set_font(subtitle.add_run("STM32F103C8T6 环境监测与 USB 串口直连电脑简化方案"), 12, color="667085")
    add_table(doc, ["项目", "确定方案"], [
        ("实验室规模", "不超过 20 平方米；现场人数不超过 5 人；2 路摄像头"),
        ("环境主控", "STM32F103C8T6 最小系统板，仅承担环境指标监测"),
        ("连接方式", "电脑 USB 直连：USB 同时提供低压供电和串口数据"),
        ("电脑平台", "既有 Flask 预警平台：现场画面、人员关联、风险事件和告警确认"),
        ("远程通知", "电脑可上网时通过飞书机器人 Webhook 发送；无网时保留本地页面告警"),
        ("编制日期", str(date.today())),
    ], [2200, 7160])
    add_callout(doc, "核心结论", "首版不需要 W5500、MQTT、网线、交换机或网络配置。使用 CP2102 或 CH340 USB 转串口模块，将 STM32 的 USART1 直接连到电脑即可。")
    add_callout(doc, "安全边界", "烟雾、火焰和可燃气体的认证报警设备必须保留独立现场报警回路。STM32 和电脑只读取辅助状态，不控制急停、安全继电器或危险负载。", warning=True)

    doc.add_heading("1. 系统范围与架构", 1)
    add_flow(doc)
    add_table(doc, ["组成", "职责"], [
        ("门禁与摄像头", "面部识别进出、现场人数与行为监测、视频画面中显示人员姓名。"),
        ("环境传感器", "温湿度、烟雾、火焰；可按已有传感器接口增加气体或液位等监测。"),
        ("STM32 最小系统板", "采样、去抖、阈值判断、心跳和串口上报；不处理视频、人脸或飞书。"),
        ("电脑预警平台", "接收串口数据，按规则产生风险预警，关联现场人员，记录事件并发送飞书通知。"),
    ], [2500, 6860])

    doc.add_heading("2. STM32 硬件与接线", 1)
    p = doc.add_paragraph()
    set_font(p.add_run("STM32F103C8T6 具备 72 MHz 主频、64 KB Flash、20 KB SRAM 和 LQFP48 引脚，足以完成本项目的传感器驱动、阈值判断和串口协议。它不承担视频分析、人脸识别、数据库或飞书接口。"))
    add_table(doc, ["功能", "STM32 引脚", "推荐接法"], [
        ("电脑串口", "PA9 / PA10", "USART1：PA9(TX) 接 USB 转串口模块 RX；模块 TX 接 PA10(RX)。"),
        ("温湿度传感器", "PB6 / PB7", "I2C1 的 SCL / SDA，适合短线数字传感器。"),
        ("RS485 传感器", "PA2 / PA3 / PB12", "USART2 TX / RX / DE-RE，经隔离 RS485 模块连接 Modbus 设备。"),
        ("烟雾 / 火焰触点", "PB10 / PB11", "先经光耦隔离或 3.3 V 电平调理，再接数字输入。"),
        ("模拟量输入", "PA0 / PA1", "仅接调理后的 0-3.0 V 信号，输入不得超过 VDDA。"),
        ("下载与调试", "PA13 / PA14", "保留 SWDIO / SWCLK；BOOT0 和 PB2/BOOT1 默认保持低电平。"),
    ], [2200, 2200, 4960], 8.8)
    add_bullets(doc, [
        "推荐购买 CP2102 或 CH340 USB 转 TTL 串口模块。选择 3.3 V 串口逻辑的版本，不能把 5 V TTL 的 TX 直接接入 PA10。",
        "接线固定为：模块 5V -> 最小系统板 5V，GND -> GND，模块 TX -> PA10，PA9 -> 模块 RX。",
        "电脑 USB 是本方案的唯一低压电源入口。禁止同时接入另一组 5 V 电源，避免反向供电。",
        "24 V、4-20 mA、长距离线缆和工业报警器不得直接接最小系统板；需要单独供电及隔离/调理。",
        "若现有最小系统板的 micro-USB 电路和 USB CDC 固件均已确认可用，可省去 USB 转串口模块；在未确认原理图前，CP2102/CH340 更稳妥。",
    ])

    doc.add_heading("3. USB 串口数据接口", 1)
    add_table(doc, ["项目", "固定约定"], [
        ("物理链路", "电脑 USB -> CP2102/CH340 -> STM32 USART1。"),
        ("串口参数", "115200 bps，8N1，UTF-8 JSON Lines（一行一条 JSON）。"),
        ("遥测周期", "每 2 至 5 秒发送一次温湿度与设备状态。"),
        ("告警上报", "烟雾、火焰、传感器故障等状态变化后立即发送，携带 sequence 和 event_id。"),
        ("设备离线", "电脑端连续 30 秒未收到有效数据时标记设备离线并产生技术告警。"),
        ("平台入口", "电脑端 serial_worker 读取 COM 口后，调用 POST /api/environment/ingest 写入现有预警服务。"),
    ], [2200, 7160])
    code = doc.add_paragraph()
    code.paragraph_format.left_indent = Inches(0.2)
    code.paragraph_format.space_after = Pt(6)
    set_font(code.add_run('{"device_id":"ENV01","sequence":1024,"temperature":26.4,"humidity":51.2,"smoke_alarm":false,"flame_alarm":false,"online":true}'), 8.5, color="344054")
    add_callout(doc, "软件改动", "现有平台已有环境模拟、告警去重、升级、确认、恢复和飞书 Webhook。硬件接入时新增串口读取服务即可，不需要 MQTT 订阅器或网络 Broker。")

    doc.add_heading("4. 实施顺序与验收", 1)
    add_numbered(doc, [
        "核对最小系统板的 5V、3.3V、GND、PA9、PA10、SWD 和 BOOT 引脚标识；确认板载稳压器可接受 USB 的 5V 输入。",
        "按现有传感器型号确认供电电压与输出接口，先完成隔离、分压或 RS485 模块，再连到 STM32。",
        "连接 CP2102/CH340 与 STM32，使用串口调试工具确认电脑能识别 COM 口并收到心跳数据。",
        "编写 STM32 采样、去抖、阈值、心跳和 JSON Lines 上报固件；电脑端实现 COM 口自动重连。",
        "将 serial_worker 接入现有 Flask 环境数据入口，联调网页告警、视频人员关联和飞书通知。",
    ])
    add_table(doc, ["验收项目", "要求"], [
        ("环境数据", "温湿度按配置周期刷新，并区分超限、异常值和传感器故障。"),
        ("硬件事件", "烟雾或火焰辅助触点变化后 1 秒内形成串口事件。"),
        ("平台联动", "电脑收到串口数据后 5 秒内在预警页面显示，并关联当前现场人员。"),
        ("USB 断开", "电脑端 30 秒内标记设备离线；重新插入 USB 后自动恢复串口通信。"),
        ("本地安全", "电脑或 STM32 故障不影响认证报警设备自身的现场声光报警。"),
        ("稳定性", "连续运行 72 小时无串口死锁、内存耗尽或异常重启。"),
    ], [2200, 7160])

    doc.add_heading("5. 采购与确认清单", 1)
    add_bullets(doc, [
        "1 个 CP2102 或 CH340 USB 转 TTL 串口模块，明确为 3.3 V 串口逻辑。",
        "1 根可靠 USB 数据线；电脑 USB 口应只为最小系统板和低功耗逻辑供电。",
        "根据传感器接口选配 I2C 模块、隔离 RS485 模块、光耦输入板或模拟量调理板。",
        "最小系统板实物正反面照片或原理图，用于最终确认 5V 输入、串口和 USB 走线。",
        "若飞书为必需通知渠道，运行预警平台的电脑需要能够访问互联网；这不改变 STM32 的无网络接入方式。",
    ])
    add_callout(doc, "资料依据", "已参考用户提供的 STM32F103C8T6 数据手册。该文件属于芯片数据手册而非具体最小系统板原理图，因此本文的板端 5V 输入需在实际板子上再核对。", warning=True)
    doc.save(OUT_PATH)
    return OUT_PATH


if __name__ == "__main__":
    print(build_document())
